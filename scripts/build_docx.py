"""Build the midterm answer sheet (DOCX) from data/results.json and figures/."""
import json, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "figures")
R = json.load(open(os.path.join(ROOT, "data", "results.json")))
D = json.load(open(os.path.join(ROOT, "data", "datasets.json")))
KEYS = sorted(R, key=lambda k: int(k[1:3]))

NAME = "Pradigta Hisyam Ramadhan"
NIM = "28726019"
DATE = "14 October 2026"
OUT = os.path.join(ROOT, f"UTS-NT6094-PradigtaHR-{NIM}.docx")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Cm(2.2))
sec.top_margin = sec.bottom_margin = Cm(2.0)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(4)
st.paragraph_format.line_spacing = 1.12
for lvl, size in ((1, 14), (2, 12)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x68)
    h.element.rPr.rFonts.set(qn("w:asciiTheme"), "")
    h.element.rPr.rFonts.set(qn("w:hAnsiTheme"), "")

fig_no = [0]
tab_no = [0]


def para(text="", bold=False, italic=False, align=None, size=None, after=None):
    p = doc.add_paragraph()
    add_runs(p, text, bold, italic, size)
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def add_runs(p, text, bold=False, italic=False, size=None):
    """Very small markup: **bold**, *italic*, _{sub}, ^{sup}."""
    import re
    tokens = re.split(r"(\*\*.+?\*\*|\*.+?\*|_\{.+?\}|\^\{.+?\})", text)
    for t in tokens:
        if not t:
            continue
        b, i, sub, sup = bold, italic, False, False
        if t.startswith("**"):
            t, b = t[2:-2], True
        elif t.startswith("*"):
            t, i = t[1:-1], True
        elif t.startswith("_{"):
            t, sub = t[2:-1], True
        elif t.startswith("^{"):
            t, sup = t[2:-1], True
        r = p.add_run(t)
        r.bold, r.italic = b, i
        r.font.subscript, r.font.superscript = sub, sup
        if size:
            r.font.size = Pt(size)
    return p


def bullet(text, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    add_runs(p, text)
    p.paragraph_format.space_after = Pt(2)
    return p


def equation(text, num):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0, c1 = t.rows[0].cells
    c0.width, c1.width = Cm(14.6), Cm(2.0)
    p = c0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, text, italic=False)
    q = c1.paragraphs[0]
    q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    q.add_run(f"({num})")


def figure(path, caption, width_cm):
    fig_no[0] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(c, f"**Figure {fig_no[0]}.** " + caption, size=9.5)
    c.paragraph_format.space_after = Pt(8)
    return fig_no[0]


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), hex_fill)
    tcPr.append(s)


def table(header, rows, caption, widths=None, size=9):
    tab_no[0] += 1
    c = doc.add_paragraph()
    add_runs(c, f"**Table {tab_no[0]}.** " + caption, size=9.5)
    c.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        cell = t.rows[0].cells[j]
        cell.text = ""
        add_runs(cell.paragraphs[0], h, bold=True, size=size)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(cell, "DCE6F2")
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ""
            add_runs(cell.paragraphs[0], str(v), size=size)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in t.rows[:-1]:
        for cell in row.cells:
            cell.paragraphs[0].paragraph_format.keep_with_next = True
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tab_no[0]


# ---------------------------------------------------------------- header
para("NT6094 Teknik Penulisan Jurnal Ilmiah — Midterm Exam (UTS) 2026-1", bold=True, size=10,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para("Feature Extraction from Ten Compression-Test Datasets: Peak, Last Point, and Linear-Region "
     "Regression as Functions of Composition g", bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
info = doc.add_table(rows=4, cols=2)
info.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (k, v) in enumerate([("Name", NAME), ("Student ID (NIM)", NIM), ("Midterm date", DATE),
                            ("Seed / datasets", "28726019 / g40s1 g42s3 g44s3 g46s4 g48s2 g50s3 g52s3 g54s4 g56s3 g58s1")]):
    info.rows[i].cells[0].width, info.rows[i].cells[1].width = Cm(4.0), Cm(12.5)
    add_runs(info.rows[i].cells[0].paragraphs[0], k, bold=True)
    add_runs(info.rows[i].cells[1].paragraphs[0], v)
para("Nanotechnology Study Program, The Graduate School of Multidisciplinary Science and Technology, Institut Teknologi Bandung. Lecturer: Sparisoma Viridi. "
     "Assignment page: https://endap.github.io/26j76/", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

# ---------------------------------------------------------------- abstract
doc.add_heading("Summary", level=1)
r40, r58 = R["g40s1"], R["g58s1"]
para("Ten datasets, selected by the seed 28726019 from the 2026-1 midterm app, were processed with the "
     "endap feature-extraction app. Each dataset is a monotonic loading curve y(x) with a toe region, a "
     "linear region, and a maximum followed by a small drop. For every dataset the peak (maximum y), the last "
     "point (largest x), and a least-squares line y = mx + b over the linear region were extracted. The linear "
     "region was chosen with one fixed rule for all datasets (points on the rising branch with 25 % to 75 % of "
     "the peak y). All fits have R^{2} ≥ 0.987. Both the peak y and the slope m fall strongly as g increases "
     f"(peak y from {r40['ypeak']:.0f} at g = 40 to {r58['ypeak']:.0f} at g = 58; m from {r40['m']:.0f} to "
     f"{r58['m']:.0f}), about 25–28 % per step of Δg = 2. The data are most consistent with compression tests on "
     "specimens whose composition is set by g: x is the crosshead displacement, y the compressive load, m the "
     "specimen stiffness, and g the share (in %) of one constituent in the mixture.")

# ---------------------------------------------------------------- 1 intro
doc.add_heading("1. Introduction", level=1)
para("The midterm page (endap post 26j76) gives every student a unique combination of ten datasets, chosen "
     "from the student ID used as a seed [1]. A dataset name such as g82s9 encodes two numbers, g = 82 and "
     "s = 9. Each dataset is published as its own endap post with an interactive app that shows the data, lets "
     "the user mark four points by their id, and computes a linear regression [2].")
para("The datasets come from a larger collection called *p2mi-wd data* [3], whose original measurement is "
     "described as a ys-measurement and whose pre-processing note lists ten compositions g40 to g58 [4]. In that "
     "note every composition has four quantities g, e, h, and t: g rises from 4.80 to 6.96, e and h fall, "
     "h = e/2, and g + e + h = t = 12.0 for all rows. The code number is therefore the share of g in percent "
     "(g40 → 0.40 × 12 = 4.80). The same note reports 40 datasets in ten clusters, which fits four repeated "
     "specimens (s = 1…4) per composition. The visualisation post that the app is built on is titled "
     "*compression-test viz* [5]. So each dataset is most likely one compression test, and the task is to "
     "extract features of the curve and see how they change with composition g.")
para("This report (i) explains how the app and the model were used, (ii) summarises peak, last point, m, and "
     "b for the ten datasets as functions of g, and (iii) discusses what x, y, and g are likely to be.")

# ---------------------------------------------------------------- 2 methods
doc.add_heading("2. Methodology", level=1)
doc.add_heading("2.1 Obtaining the datasets", level=2)
para("The midterm page was opened and the seed 28726019 (the student ID) was selected. The app then shows "
     "the dataset names and the links to their posts (Figure 1). Each link was opened in a new browser tab and "
     "the post title was checked against the dataset name; all ten titles match (Table 1).")
figure(os.path.join(FIG, "app_00_seed.png"),
       "Screenshot of the midterm app (https://endap.github.io/26j76/) with seed 28726019 selected, showing "
       "the ten dataset names and their links.", 15.5)
rows = []
for k in KEYS:
    v = D[k]
    rows.append([k, k[1:3], k[4:], v["link"], v["title"], v["n"], v["every"], len(v["ds"])])
table(["Dataset", "g", "s", "Link", "Post title", "Raw rows", "every", "Points (id)"], rows,
      "Datasets for seed 28726019. Post titles confirm the dataset names. *every* is the sampling step "
      "used by the app on that page; *Points* is the number of selectable ids.",
      widths=[1.7, 1.0, 0.8, 1.6, 2.8, 1.7, 1.3, 1.9])

doc.add_heading("2.2 How the app works", level=2)
para("Each dataset post holds the raw data as two columns. The app takes the **second column as x** "
     "(independent variable) and the **first column as y** (dependent variable). It keeps every *every*-th row "
     "(10, 20, or 25 depending on the page, Table 1) and numbers these points with an id 0000, 0001, … shown "
     "in the console panel together with the original row number. The UI has three parts: the chart (top), the "
     "data console (lower left), and the input/output panel with four buttons (lower right) [2]. The steps are:")
for s in ["**load** writes a default input with four keys: PEAK-ID, LAST-ID, LINI-ID, LINF-ID.",
          "The ids are edited: PEAK-ID = point with the maximum y; LAST-ID = point with the last (largest) x; "
          "LINI-ID and LINF-ID = initial and final point of the linear regression.",
          "**read** places the four points on the chart so the choice can be checked visually; this is repeated "
          "until the points are where they should be.",
          "**calc** fits y = mx + b to all points from LINI-ID to LINF-ID and prints m, b, and R^{2}; the "
          "model line is drawn from id LINI−5 to LINF+15 (so LINI ≥ 5 and LINF ≤ last id − 15 are required).",
          "**clear** resets input, output, and chart."]:
    bullet(s, "List Number")

doc.add_heading("2.3 Model", level=2)
para("A straight line is fitted to the linear part of each curve [5], [6]:")
equation("y = m x + b", 1)
para("with the least-squares slope and intercept over the n selected points [6]")
equation("m = [ n Σxy − Σx Σy ] / [ n Σx^{2} − (Σx)^{2} ]", 2)
equation("b = ( Σy − m Σx ) / n", 3)
para("and the coefficient of determination [7]")
equation("R^{2} = [ n Σxy − Σx Σy ]^{2} / { [ n Σx^{2} − (Σx)^{2} ] [ n Σy^{2} − (Σy)^{2} ] }", 4)
para("R^{2} close to 1 means the selected points lie close to a straight line. The x-axis intercept "
     "x_{0} = −b/m is also used in the discussion because it removes the effect of the toe region.")

doc.add_heading("2.4 Rule for choosing the ids", level=2)
para("Choosing the linear region by eye differs between persons and between datasets. To make the ten results "
     "comparable, one rule was used for all datasets:")
for s in ["PEAK-ID: the id with the largest y.",
          "LAST-ID: the last id (largest x).",
          "LINI-ID: the first id on the rising branch (before the peak) with y ≥ 0.25 y_{peak}.",
          "LINF-ID: the last id before the peak with y ≤ 0.75 y_{peak}."]:
    bullet(s)
para("This band avoids the curved toe at the start (seating of the specimen) and the rounding near the "
     "maximum. Bands of 10–90 %, 20–80 %, and 30–70 % were also tried; m changed by only a few percent and the "
     "trend with g was the same, so 25–75 % was kept. The chosen ids were entered in each app and the chart "
     "was checked after **read** (screenshots in Appendix A).")

doc.add_heading("2.5 Verification", level=2)
para("The raw data of all ten posts were also processed with a short Python script that copies the app's "
     "sampling and Eqs. (2)–(4). The values of m, b, and R^{2} returned by the app (**calc**) and by the script "
     "are identical to all printed digits for every dataset. The script and the extracted values are kept with "
     "the answer sheet and published at https://github.com/pradigtahr/nt6094-uts-2026-1 (scripts/, data/results.json).")

doc.add_heading("2.6 Flowchart", level=2)
para("Figure 2 shows the full process, from the seed to the answer sheet. It is written in Mermaid with "
     "on-page connectors, following the course material [8]; the source code is given in Appendix B.")
figure(os.path.join(FIG, "flowchart.png"),
       "Flowchart of how the datasets were obtained, processed in the app, and analysed. "
       "Green circles are on-page connectors.", 12.5)

# ---------------------------------------------------------------- 3 results
doc.add_heading("3. Results and Discussion", level=1)
doc.add_heading("3.1 Extracted values", level=2)
para("Table 2 lists the ids used in the app. Table 3 summarises the ten values of peak, last point, m, and b "
     "as functions of g.")
rows = [[k[1:3], k, f"{R[k]['peak']:04d}", f"{R[k]['last']:04d}", f"{R[k]['lini']:04d}",
         f"{R[k]['linf']:04d}", R[k]["linf"] - R[k]["lini"] + 1] for k in KEYS]
table(["g", "Dataset", "PEAK-ID", "LAST-ID", "LINI-ID", "LINF-ID", "n (fit)"], rows,
      "Point ids entered in the app for each dataset (rule of Section 2.4).",
      widths=[1.0, 1.8, 1.8, 1.8, 1.8, 1.8, 1.6])
rows = [[k[1:3], k, f"{R[k]['xpeak']:.4f}", f"{R[k]['ypeak']:.2f}", f"{R[k]['xlast']:.4f}",
         f"{R[k]['ylast']:.2f}", f"{R[k]['m']:.2f}", f"{R[k]['b']:.2f}", f"{R[k]['R2']:.5f}"] for k in KEYS]
table(["g", "Dataset", "x_{peak}", "y_{peak}", "x_{last}", "y_{last}", "m", "b", "R^{2}"], rows,
      "Peak, last point, and linear-regression parameters m and b as functions of g.",
      widths=[0.9, 1.6, 1.6, 1.7, 1.6, 1.7, 1.9, 2.1, 1.7])

f_res = figure(os.path.join(FIG, "results_vs_g.png"),
               "xy-scatter charts of the extracted values as functions of g: (a) y of peak and last point, "
               "(b) x of peak and last point, (c) slope m, (d) intercept b.", 15.0)
f_all = figure(os.path.join(FIG, "all_datasets.png"),
               "All ten raw curves y(x); darker lines have larger g.", 14.5)
f_fit = figure(os.path.join(FIG, "fits_all.png"),
               "Each dataset with the points used in the regression (dark blue), the fitted line (dashed), "
               "the peak (triangle), and the last point (square).", 15.0)

doc.add_heading("3.2 Shape of the curves", level=2)
para(f"All curves have the same three parts (Figures {f_all} and {f_fit}): (1) a flat toe where y stays near "
     "zero while x grows, (2) a long, almost straight rise, and (3) a rounded top that ends at a maximum, "
     "after which y drops slightly until the last point. The drop from peak to last point is 1.7–12 % "
     f"(largest for g44s3, {100*(R['g44s3']['ypeak']-R['g44s3']['ylast'])/R['g44s3']['ypeak']:.0f} %). "
     "The length of the toe differs strongly from test to test (x_{0} = −b/m from 0.59 to 3.82) and does not "
     "follow g. A toe of this kind is typical for the start of a mechanical test, when the loading plate first "
     "closes gaps and seats on an uneven specimen surface. g44s3 also shows a sudden drop of y at x ≈ 3.65 "
     "before the main rise, which looks like slipping or re-seating of the specimen. Datasets with large g "
     "(g48 and higher) carry a regular small ripple on top of the curve; it is the reason why g52s3 and g56s3 "
     "have the lowest R^{2} (0.988 and 0.990), but the fits still follow the data well.")

doc.add_heading("3.3 Trends with g", level=2)
ks = KEYS
para(f"**Peak and last y.** y_{{peak}} is highest at g = 44 ({R['g44s3']['ypeak']:.0f}) and falls to "
     f"{R['g58s1']['ypeak']:.0f} at g = 58 (Figure {f_res}a). A log-linear fit gives a decrease of about 25 % "
     "per Δg = 2 (Pearson r = −0.89 between g and y_{peak}). The last point follows the peak closely, so "
     "both describe the same property: the maximum load the specimen can carry.")
para(f"**Slope m.** m falls from {R['g40s1']['m']:.0f} at g = 40 to about 400 at g = 56–58 "
     f"(Figure {f_res}c), roughly 28 % per Δg = 2 (r = −0.89). Specimens with more g are therefore both "
     "weaker and softer. Two points break the smooth trend: g46s4 (m higher than g44s3) and g52s3 (m higher "
     "than g50s3). Both have the shortest curves and short toes, and only one specimen per composition was "
     "given, so the scatter between repeated specimens (s) cannot be separated from the effect of g here.")
para(f"**Intercept b.** b is negative for all datasets and moves towards zero as g rises (Figure {f_res}d, "
     "r = +0.85). b alone has little physical meaning: it is set by the slope and by where the linear part "
     "starts, b = −m x_{0}. Its trend mostly mirrors the fall of m and the varying toe length.")
para(f"**x of peak and last point.** x_{{peak}} and x_{{last}} jump between 1.3 and 5.8 with no clear trend "
     f"(Figure {f_res}b), because they include the random toe length. The part after the toe, "
     "x_{peak} − x_{0}, is 0.66–2.54 and tends to be larger for higher g, which agrees with a softer, more "
     "ductile material.")

doc.add_heading("3.4 What x, y, and g might be", level=2)
for s in [
    "**x — displacement (deformation) of the specimen**, most likely the crosshead travel in mm. It is the "
    "column that rises steadily in time, it is used as the independent variable, it includes a seating toe, "
    "and its values (0–6) fit the shortening of a small specimen in millimetres.",
    "**y — compressive load (force)**, most likely in newtons. It rises with x, reaches a maximum, and then "
    "drops when the specimen starts to fail. The maximum y_{peak} is then the failure load; dividing by the "
    "loaded area would give the compressive strength. The slope m = dy/dx is the stiffness of the specimen "
    "(N/mm); with the specimen geometry it would give the elastic (compressive) modulus.",
    "**g — composition: the share of one constituent in the mixture, in percent.** The pre-processing note "
    "gives g = (code/100) × 12 for a fixed total t = 12.0 (e.g. grams), while the other two constituents e and "
    "h share the rest in a 2 : 1 ratio [4]. Increasing g lowers both strength and stiffness, so this "
    "constituent is likely a weaker or softer phase (for example a filler, binder, or liquid fraction) — the "
    "exact material is not stated in the sources.",
    "**s — specimen (replicate) number** within one composition. Forty datasets in ten composition clusters "
    "[4] means four specimens per g; the seed gives one of them (s = 1–4) for each g."]:
    bullet(s)
para("Because the material, units, and specimen size are not given, these meanings are reasoned "
     "interpretations of the data and of the source notes, not confirmed facts.")

doc.add_heading("3.5 Limitations", level=2)
for s in ["Only one specimen per composition is analysed, so specimen-to-specimen scatter is unknown.",
          "The app works on sampled points (every 10, 20, or 25 rows); the peak of the raw data can be "
          "slightly higher than the sampled peak (e.g. g42s3: 6287 raw vs 6286 sampled).",
          "The 25–75 % band is a convention; a different band changes m by a few percent but not the trend.",
          "m is a stiffness (load per displacement), not a modulus, because specimen size is unknown."]:
    bullet(s)

# ---------------------------------------------------------------- 4 conclusions
doc.add_heading("4. Conclusions", level=1)
for s in [
    "Ten datasets (g40s1 … g58s1) were obtained with seed 28726019 and processed in the endap app. Peak, last "
    "point, m, b, and R^{2} were extracted with one fixed rule; app and independent script agree exactly.",
    "Each curve is a compression loading curve: toe, linear rise (all R^{2} ≥ 0.987), maximum, small drop.",
    "Peak load and slope m both decrease strongly as g rises (r ≈ −0.89; ~25–28 % per Δg = 2): more of "
    "constituent g gives a weaker and softer specimen.",
    "x is most likely displacement, y compressive load, g the percentage of one constituent in a 12-unit "
    "mixture, and s the specimen number. The x-values and b are dominated by the seating toe and should be "
    "corrected (x − x_{0}) before comparing deformation between tests.",
    "Repeated specimens per composition and the specimen geometry are needed to turn these features into "
    "material properties (strength, modulus) with uncertainties."]:
    bullet(s)

# ---------------------------------------------------------------- 5 AI declaration
doc.add_heading("5. Declaration of AI Use", level=1)
para("I declare that I used a generative AI tool in this midterm: **Claude Code (Claude Opus 5.5, "
     "Anthropic)**, on 14 October 2026. It was used to read the assignment instructions, download the ten "
     "dataset posts, write the Python scripts that reproduce the app's calculation, propose the id-selection "
     "rule, operate the app in the browser to take the screenshots, draw the charts and the Mermaid "
     "flowchart, draft the text of this answer sheet, and prepare the public GitHub repository with the "
     "exported chat log. I chose the datasets with my own student ID, "
     "checked the app outputs against the script outputs, and reviewed the interpretation. I am responsible "
     "for the content of this answer sheet.")

# ---------------------------------------------------------------- 6 references
doc.add_heading("6. References and AI Chats", level=1)
refs = [
    "S. Viridi, \"2026-1 midterm-datasets\", endap, 13 Oct 2026, url https://endap.github.io/26j76/ [20261014].",
    "S. Viridi, \"g40s1-data\" … \"g58s1-data\" (posts 26j34, 26j40, 26j44, 26j49, 26j51, 26j56, 26j60, "
    "26j65, 26j68, 26j70), endap, 12 Oct 2026, url https://endap.github.io/26j34/ etc. [20261014].",
    "S. Viridi, \"p2mi-wd data\", endap, 9 Oct 2026, url https://endap.github.io/26j27/ [20261014].",
    "S. Viridi, \"pre-proc p2mi-wd yessi-data\", endap, 9 Oct 2026, url https://endap.github.io/26j26/ [20261014].",
    "S. Viridi, \"compression-test viz\", endap, 11 Oct 2026, url https://endap.github.io/26j32/ [20261014].",
    "L. Treible, M. Szulczewski, \"How Do I Calculate a Linear Regression? Finding Relationships between "
    "Variables in the Earth Sciences\", Math You Need - Majors, 11 Aug 2023, url https://serc.carleton.edu/277935 [20261014].",
    "M. Gupta, \"R Squared | Coefficient of Determination\", GeeksforGeeks, 20 May 2026, url "
    "https://www.geeksforgeeks.org/maths/r-squared/ [20261014].",
    "S. Viridi, \"mermaid diagram hands-on\", endap, 30 Sep 2026, url https://endap.github.io/26i75/ [20261014].",
    "W. D. Callister, D. G. Rethwisch, *Materials Science and Engineering: An Introduction*, 10th ed., "
    "Wiley, 2018 (Ch. 6, Mechanical properties: stress–strain behaviour, elastic modulus, strength).",
    "ASTM D695-15, \"Standard Test Method for Compressive Properties of Rigid Plastics\", ASTM International, "
    "2015 (seating/toe compensation of the load–displacement curve).",
    "Claude Code (Claude Opus 5.5, Anthropic), \"NT6094 midterm (UTS) 2026-1: processing the datasets of seed "
    "28726019\", chat log, 14 Oct 2026, url https://github.com/pradigtahr/nt6094-uts-2026-1/blob/main/ai-chat/claude-code-chat-log.md [20261014].",
]
for i, r in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)
    p.paragraph_format.space_after = Pt(3)
    add_runs(p, f"[{i}]\t" + r, size=10)

# ---------------------------------------------------------------- appendix
doc.add_page_break()
doc.add_heading("Appendix A. App screenshots after calc", level=1)
para("Each screenshot shows the post title (dataset name), the chart with peak, last, linear and model points, "
     "the entered ids, and the m, b, R^{2} output.")
grid = doc.add_table(rows=len(KEYS) // 2, cols=2)
grid.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, k in enumerate(KEYS):
    fig_no[0] += 1
    cell = grid.rows[i // 2].cells[i % 2]
    cell.width = Cm(8.3)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(FIG, f"app_{k}.png"), width=Cm(7.9))
    c = cell.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(c, f"**Figure {fig_no[0]}.** App for {k} ({D[k]['link']}): PEAK-ID {R[k]['peak']:04d}, "
                f"LAST-ID {R[k]['last']:04d}, LINI-ID {R[k]['lini']:04d}, LINF-ID {R[k]['linf']:04d}.", size=8.5)
    c.paragraph_format.space_after = Pt(6)
doc.add_heading("Appendix B. Mermaid code of the flowchart", level=1)
code = open(os.path.join(FIG, "flowchart.mmd"), encoding="utf-8").read()
for line in code.rstrip().split("\n"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(line if line else " ")
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)

doc.save(OUT)
print(OUT)
